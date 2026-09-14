import csv
import io
from decimal import Decimal, InvalidOperation

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from ..core.db import get_db
from ..models import ChangeLog, InventoryBalance, Product, ProductVariation, User
from ..schemas import ProductImportCommit, ProductImportRow
from .deps import current_admin
from .pos import product_payload

router = APIRouter(prefix="/api/v1/products", tags=["products"])
FIELDS = ["sku", "name", "variation_name", "variation_sku", "selling_price", "initial_stock", "category", "enable_stock", "is_active"]


def parse_bool(value, default=True):
    if value in (None, ""):
        return default
    return str(value).strip().lower() in {"1", "true", "yes", "ya", "aktif"}


def parse_rows(filename: str, content: bytes) -> tuple[list[dict], list[dict]]:
    if filename.lower().endswith(".xlsx"):
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        values = list(sheet.iter_rows(values_only=True))
        headers = [str(value or "").strip() for value in values[0]] if values else []
        raw_rows = [dict(zip(headers, row)) for row in values[1:]]
    else:
        text = content.decode("utf-8-sig")
        raw_rows = list(csv.DictReader(io.StringIO(text)))
    valid, errors = [], []
    seen = set()
    for number, row in enumerate(raw_rows, 2):
        try:
            sku = str(row.get("sku") or "").strip()
            name = str(row.get("name") or "").strip()
            variation_name = str(row.get("variation_name") or "DUMMY").strip() or "DUMMY"
            variation_sku = str(row.get("variation_sku") or sku).strip()
            if not sku or not name or not variation_sku:
                raise ValueError("sku, name, dan variation_sku wajib diisi")
            if variation_sku.lower() in seen:
                raise ValueError("variation_sku duplikat di dalam file")
            seen.add(variation_sku.lower())
            parsed = ProductImportRow(
                sku=sku, name=name, variation_name=variation_name, variation_sku=variation_sku,
                selling_price=Decimal(str(row.get("selling_price") or 0)),
                initial_stock=Decimal(str(row.get("initial_stock") or 0)),
                category=str(row.get("category") or "").strip() or None,
                enable_stock=parse_bool(row.get("enable_stock")), is_active=parse_bool(row.get("is_active")),
            )
            valid.append(parsed.model_dump(mode="json"))
        except (ValueError, InvalidOperation) as exc:
            errors.append({"row": number, "message": str(exc)})
    return valid, errors


@router.post("/import/preview")
async def preview(file: UploadFile = File(...), user: User = Depends(current_admin)):
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(413, "Ukuran file maksimum 10 MB.")
    if not file.filename or not file.filename.lower().endswith((".csv", ".xlsx")):
        raise HTTPException(422, "Gunakan file CSV atau XLSX.")
    rows, errors = parse_rows(file.filename, content)
    return {"data": {"rows": rows, "errors": errors, "valid_count": len(rows), "error_count": len(errors)}}


@router.post("/import/commit")
def commit(payload: ProductImportCommit, user: User = Depends(current_admin), db: Session = Depends(get_db)):
    from ..models import Location
    if not db.scalar(select(Location.id).where(Location.id == payload.location_id, Location.business_id == user.business_id)):
        raise HTTPException(422, "Lokasi import tidak ditemukan.")
    imported = 0
    for row in payload.rows:
        product = db.scalar(select(Product).where(Product.business_id == user.business_id, Product.sku == row.sku))
        if not product:
            product = Product(business_id=user.business_id, sku=row.sku, name=row.name)
            db.add(product)
            db.flush()
        product.name, product.category = row.name, row.category
        product.enable_stock, product.is_active = row.enable_stock, row.is_active
        product.revision += 1
        variation_sku = row.variation_sku or row.sku
        variation = db.scalar(select(ProductVariation).where(ProductVariation.business_id == user.business_id, ProductVariation.sku == variation_sku))
        if variation and variation.product_id != product.id:
            raise HTTPException(409, f"variation_sku {variation_sku} sudah digunakan produk lain.")
        if not variation:
            variation = ProductVariation(business_id=user.business_id, product=product, sku=variation_sku)
            db.add(variation)
            db.flush()
        variation.name, variation.selling_price = row.variation_name, row.selling_price
        variation.revision += 1
        balance = db.scalar(select(InventoryBalance).where(InventoryBalance.location_id == payload.location_id, InventoryBalance.variation_id == variation.id))
        if not balance:
            balance = InventoryBalance(business_id=user.business_id, location_id=payload.location_id, variation=variation)
            db.add(balance)
        balance.quantity = row.initial_stock
        balance.revision = (balance.revision or 0) + 1
        db.flush()
        data = product_payload(product, payload.location_id)
        db.add(ChangeLog(business_id=user.business_id, entity_type="product", entity_uuid=product.uuid, revision=product.revision, payload=data))
        imported += 1
    db.commit()
    return {"data": {"imported": imported}}


@router.get("/export")
def export_products(user: User = Depends(current_admin), db: Session = Depends(get_db)):
    products = db.scalars(select(Product).options(joinedload(Product.variations).joinedload(ProductVariation.balances)).where(Product.business_id == user.business_id)).unique().all()
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=FIELDS)
    writer.writeheader()
    for product in products:
        for variation in product.variations:
            stock = sum((balance.quantity for balance in variation.balances), Decimal("0"))
            writer.writerow({
                "sku": product.sku, "name": product.name, "variation_name": variation.name,
                "variation_sku": variation.sku, "selling_price": variation.selling_price,
                "initial_stock": stock, "category": product.category or "",
                "enable_stock": int(product.enable_stock), "is_active": int(product.is_active),
            })
    body = io.BytesIO(output.getvalue().encode("utf-8-sig"))
    return StreamingResponse(body, media_type="text/csv", headers={"Content-Disposition": "attachment; filename=produk-pos.csv"})
